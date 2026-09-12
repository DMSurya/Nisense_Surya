"""Multi-sheet Excel export of `readings`, with device/date/type filters.

Mirrors the mobile app's `RecordXlsxExporter` sheet layout (Summary, Glucose,
Glucose_Raw, PPG_Raw, Vitals, Temp) so the same file structure works whether
it came off the phone directly or was exported here after a cloud push.
"""

from __future__ import annotations

import datetime as dt
import io

from fastapi import APIRouter, Depends
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from openpyxl.worksheet.worksheet import Worksheet
from sqlalchemy import select
from sqlalchemy.orm import Session

from ..db import get_db
from ..models import Reading
from ..security import get_current_principal

router = APIRouter(tags=["export"])

_XLSX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _header(sheet: Worksheet, cols: list[str]) -> None:
    sheet.append(cols)


def _write_glucose(sheet: Worksheet, rows: list[Reading]) -> None:
    _header(sheet, [
        "Record_ID", "Timestamp", "Patient_ID", "Device_ID",
        "Glucose_mg_dL", "Quality", "Variant", "Model_Version",
        "Intercept", "Outlier_K", "Tot_Coeff", "Y1", "Avg", "StdDev", "UpLim", "LlLim",
        "P_Count", "N_Count", "P_Val", "N_Val", "P_Plus_N",
        "Y2_Val", "Y2_Percent", "Group_CD", "Y2_Factor", "Y2_Factor_Val",
        "Const_Val", "Y3_Value", "Y3_Row", "Elim_Per", "Elim_Val", "Y_Value",
        "Cal_Factor", "AG_Adj", "Norm_Glucose", "Insulin", "Insulin_Corr",
        "Insulin_Ratio", "Inv_Ratio", "HOMA_IR",
    ])
    for r in rows:
        e = r.extra or {}
        sheet.append([
            r.record_id, r.ts.isoformat() if r.ts else None, r.patient_id, r.device_id,
            e.get("glucose_mg_dl"), e.get("quality"), e.get("variant"), e.get("model_version"),
            e.get("intercept"), e.get("outlier_k"), e.get("tot_coeff"), e.get("y1_value"),
            e.get("avg_val"), e.get("std_dev"), e.get("up_lim"), e.get("ll_lim"),
            e.get("p_count"), e.get("n_count"), e.get("p_val"), e.get("n_val"), e.get("p_plus_n"),
            e.get("y2_val"), e.get("y2_percent"), e.get("group_cd"), e.get("y2_factor"),
            e.get("y2_factor_val"), e.get("const_val"), e.get("y3_value"), e.get("y3_row_no"),
            e.get("elim_per"), e.get("elim_val"), e.get("y_value"), e.get("calibration_factor"),
            e.get("ag_adjusted"), e.get("normalized_glucose"), e.get("actual_insulin"),
            e.get("insulin_correction"), e.get("insulin_ratio"), e.get("inverse_ratio"),
            e.get("homa_ir_index"),
        ])


def _write_vitals(sheet: Worksheet, rows: list[Reading]) -> None:
    _header(sheet, [
        "Record_ID", "Timestamp", "Patient_ID", "Device_ID",
        "HR_BPM", "HR_Conf", "SpO2_Pct", "SpO2_Conf",
        "Hb_g_dL", "Hb_Conf", "Resp_BPM", "Resp_Conf",
        "SDNN_ms", "RMSSD_ms", "Sys_mmHg", "Dia_mmHg",
        "Quality", "SNR_dB_x10", "Perf_x10", "Sample_Rate_Hz", "Sample_Count",
    ])
    for r in rows:
        e = r.extra or {}
        sheet.append([
            r.record_id, r.ts.isoformat() if r.ts else None, r.patient_id, r.device_id,
            e.get("hr_bpm"), e.get("hr_conf"), e.get("spo2_percent"), e.get("spo2_conf"),
            e.get("hb_g_dl"), e.get("hb_conf"), e.get("resp_rate_bpm"), e.get("resp_conf"),
            e.get("sdnn_ms"), e.get("rmssd_ms"), e.get("systolic_mmhg"), e.get("diastolic_mmhg"),
            e.get("quality"), e.get("snr_db_x10"), e.get("perfusion_index_x10"),
            e.get("sample_rate_hz"), e.get("sample_count"),
        ])


def _write_temp(sheet: Worksheet, rows: list[Reading]) -> None:
    _header(sheet, [
        "Record_ID", "Timestamp", "Patient_ID", "Device_ID",
        "SoC_Temp_C", "Skin_Temp_C", "Skin_Band", "Source",
    ])
    for r in rows:
        e = r.extra or {}
        sheet.append([
            r.record_id, r.ts.isoformat() if r.ts else None, r.patient_id, r.device_id,
            e.get("soc_temp_c"), e.get("skin_temp_c"), e.get("skin_band"), e.get("source"),
        ])


def _write_ppg_raw(sheet: Worksheet, rows: list[Reading]) -> None:
    _header(sheet, [
        "Parent_ID", "Chunk_Index", "Sample_Index",
        "IR", "Red", "Green", "IR_DC", "Red_DC", "Green_DC",
        "IR_AC", "Red_AC", "Green_AC", "Accel_X", "Accel_Y", "Accel_Z",
        "Device_ID", "Patient_ID",
    ])
    for r in rows:
        e = r.extra or {}
        header = e.get("header") or {}
        for idx, s in enumerate(e.get("samples") or []):
            sheet.append([
                e.get("parent_id"), header.get("chunk_index"), idx,
                s.get("ir"), s.get("red"), s.get("green"),
                s.get("ir_dc"), s.get("red_dc"), s.get("green_dc"),
                s.get("ir_ac"), s.get("red_ac"), s.get("green_ac"),
                s.get("accel_x"), s.get("accel_y"), s.get("accel_z"),
                r.device_id, r.patient_id,
            ])


def _write_glucose_raw(sheet: Worksheet, rows: list[Reading]) -> None:
    _header(sheet, [
        "Parent_ID", "Chunk_Index", "Sample_Index", "ADC", "Voltage_mV",
        "Device_ID", "Patient_ID",
    ])
    for r in rows:
        e = r.extra or {}
        header = e.get("header") or {}
        for idx, s in enumerate(e.get("samples") or []):
            sheet.append([
                e.get("parent_id"), header.get("chunk_index"), idx,
                s.get("adc"), s.get("voltage_mv"),
                r.device_id, r.patient_id,
            ])


def _write_other(sheet: Worksheet, rows: list[Reading]) -> None:
    _header(sheet, [
        "ID", "Timestamp", "Device_ID", "Patient_ID", "Type", "Record_ID", "Value", "Quality",
    ])
    for r in rows:
        sheet.append([
            r.id, r.ts.isoformat() if r.ts else None, r.device_id, r.patient_id,
            r.type, r.record_id, r.value, r.quality,
        ])


def _write_summary(
    sheet: Worksheet,
    by_type: dict[str, list[Reading]],
    total: int,
    filters: dict[str, object],
) -> None:
    _header(sheet, ["Field", "Value"])
    sheet.append(["Export_Date", dt.datetime.now(dt.timezone.utc).isoformat()])
    for k, v in filters.items():
        if v is not None:
            sheet.append([k, str(v)])
    sheet.append(["", ""])
    sheet.append(["Record_Type", "Count"])
    for t, rows in sorted(by_type.items()):
        sheet.append([t, len(rows)])
    sheet.append(["Total", total])


_SHEET_WRITERS = {
    "glucose": ("Glucose", _write_glucose),
    "vitals": ("Vitals", _write_vitals),
    "temp": ("Temp", _write_temp),
    "ppg_raw": ("PPG_Raw", _write_ppg_raw),
    "glucose_raw": ("Glucose_Raw", _write_glucose_raw),
}


@router.get("/export/readings.xlsx", dependencies=[Depends(get_current_principal)])
def export_readings_xlsx(
    db: Session = Depends(get_db),
    device_id: str | None = None,
    patient_id: str | None = None,
    type: str | None = None,  # noqa: A002 - matches query param name used elsewhere
    since: dt.datetime | None = None,
    until: dt.datetime | None = None,
) -> StreamingResponse:
    stmt = select(Reading)
    if device_id:
        stmt = stmt.where(Reading.device_id == device_id)
    if patient_id:
        stmt = stmt.where(Reading.patient_id == patient_id)
    if type:
        stmt = stmt.where(Reading.type == type)
    if since:
        stmt = stmt.where(Reading.ts >= since)
    if until:
        stmt = stmt.where(Reading.ts <= until)
    stmt = stmt.order_by(Reading.ts.asc())
    rows = list(db.scalars(stmt).all())

    by_type: dict[str, list[Reading]] = {}
    for r in rows:
        by_type.setdefault(r.type, []).append(r)

    wb = Workbook()
    summary_sheet = wb.active
    summary_sheet.title = "Summary"
    _write_summary(
        summary_sheet,
        by_type,
        len(rows),
        {
            "device_id": device_id,
            "patient_id": patient_id,
            "type": type,
            "since": since,
            "until": until,
        },
    )

    for t, type_rows in by_type.items():
        sheet_name, writer = _SHEET_WRITERS.get(t, (None, None))
        if sheet_name is None:
            sheet = wb.create_sheet(f"Other_{t}"[:31])
            _write_other(sheet, type_rows)
        else:
            writer(wb.create_sheet(sheet_name), type_rows)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)

    ts = dt.datetime.now().strftime("%Y_%m_%d-%H_%M_%S")
    filename = f"nisense_export_{ts}.xlsx"
    return StreamingResponse(
        buf,
        media_type=_XLSX_MEDIA_TYPE,
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )
